"""错题本导出：JSON 模板 + 压缩图片 + 静态资源目录，打包成一个 ZIP。

为什么不走 PDF：PDF 里图片与排版都要自己算，格式问题最多。这里先把内容收成一份
标准 JSON（模板 + 数据），再由它渲染出 docx 与可直接打开的 html；题目图片统一压缩
进 assets/ 并用相对路径引用，所以这个包换台机器打开也不会丢图。
"""
from __future__ import annotations

import io
import json
import os
import re
import tempfile
import zipfile
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

ASSET_DIR = "assets"
IMAGE_MAX_EDGE = 1200
IMAGE_QUALITY = 78

# 模板里可替换 / 循环的位置。客户要的是「前端传 JSON 模板，body 里做循环」，
# 这里把它收敛成几个明确字段：前端传什么都不至于把渲染打挂。
DEFAULT_TEMPLATE: Dict[str, Any] = {
    "header": "{title} · 错题本",
    "subheader": "{student_name}（{student_no}）· 错题 {wrong_count} 题",
    "footer": "由 AI 伴学协同平台导出 · {exported_at}",
    "section_title": "第 {index} 题 · {type_label} · 得分 {score}/{max_score}",
    "fields": ["stem", "my_answer", "correct_answer", "analysis"],
}

FIELD_LABELS = {
    "stem": "题目",
    "my_answer": "我的答案",
    "correct_answer": "正确答案",
    "analysis": "解析",
}


def merge_template(raw: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """把前端传进来的模板收敛到已知字段，未知/非法一律忽略，用默认值兜底。"""
    merged = dict(DEFAULT_TEMPLATE)
    if not isinstance(raw, dict):
        return merged
    for key in ("header", "subheader", "footer", "section_title"):
        value = raw.get(key)
        if isinstance(value, str) and value.strip():
            merged[key] = value
    fields = raw.get("fields")
    if isinstance(fields, list):
        picked = [f for f in fields if f in FIELD_LABELS]
        if picked:
            merged["fields"] = picked
    return merged


def _fmt(template: str, values: Dict[str, Any]) -> str:
    """只替换已知占位符，缺值留空。不用 str.format —— 题干里的花括号会把渲染打挂。"""

    def repl(match: "re.Match[str]") -> str:
        value = values.get(match.group(1))
        return "" if value is None else str(value)

    return re.sub(r"\{(\w+)\}", repl, template or "")


def _num(value: Any) -> str:
    if value is None:
        return "—"
    number = float(value)
    return str(int(number)) if number.is_integer() else f"{number:.1f}"


def _escape(text: Any) -> str:
    return str(text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _plain(question: Dict[str, Any], field: str) -> str:
    if field == "stem":
        return question.get("stem") or ""
    if field == "my_answer":
        return question.get("my_answer") or ""
    if field == "correct_answer":
        return question.get("correct_answer") or ""
    if field == "analysis":
        return question.get("analysis") or ""
    return ""


def _question_values(question: Dict[str, Any], index: int) -> Dict[str, Any]:
    return {
        "index": index,
        "type_label": question.get("type_label") or "",
        "score": _num(question.get("score")),
        "max_score": _num(question.get("max_score")),
    }


def compress_image(raw: bytes, index: int) -> Tuple[str, bytes]:
    """统一转 JPEG 并限长边，体积可控。Pillow 已在依赖里（fpdf2 带进来的）。"""
    from PIL import Image

    with Image.open(io.BytesIO(raw)) as img:
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        img.thumbnail((IMAGE_MAX_EDGE, IMAGE_MAX_EDGE))
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=IMAGE_QUALITY, optimize=True)
    return f"image-{index}.jpg", buffer.getvalue()


def prepare_assets(notebook: Dict[str, Any]) -> List[Tuple[str, bytes]]:
    """把题目图片压进 assets/，并把引用改写成相对路径。返回 (相对路径, 内容) 列表。

    单张图取不到或解不开都不该让整份错题本导出失败：降级成保留原始 URL 并标 missing。
    """
    from app.services.minio_service import MinioService

    assets: List[Tuple[str, bytes]] = []
    index = 0
    for question in notebook.get("questions") or []:
        rewritten: List[Dict[str, Any]] = []
        for image in question.get("images") or []:
            object_name = image.get("object_name")
            if not object_name:
                continue
            try:
                raw = MinioService.download_file(object_name)
                filename, payload = compress_image(raw, index)
            except Exception:
                rewritten.append({"url": image.get("url") or "", "missing": True})
                continue
            index += 1
            assets.append((f"{ASSET_DIR}/{filename}", payload))
            rewritten.append({"asset": f"{ASSET_DIR}/{filename}", "missing": False})
        question["images"] = rewritten
    return assets


def render_html(notebook: Dict[str, Any], template: Dict[str, Any]) -> str:
    values = dict(notebook.get("meta") or {})
    parts = [
        "<!DOCTYPE html>",
        '<html lang="zh-CN"><head><meta charset="utf-8">',
        f"<title>{_escape(values.get('title') or '错题本')}</title>",
        "<style>",
        "body{font-family:system-ui,'Microsoft YaHei',sans-serif;max-width:840px;margin:0 auto;padding:24px;line-height:1.75;color:#1f2937}",
        "h1{font-size:20px}h2{font-size:15px;margin-top:28px;border-left:4px solid #2563eb;padding-left:8px}",
        ".meta{color:#6b7280;font-size:13px}.label{font-weight:600;color:#374151;margin-top:10px}",
        ".ans{white-space:pre-wrap;background:#f9fafb;padding:8px 10px;border-radius:6px}",
        ".mine{background:#fef2f2}.right{background:#f0fdf4}",
        "img{max-width:100%;margin-top:6px;border-radius:4px}",
        "footer{margin-top:32px;color:#9ca3af;font-size:12px}",
        "</style></head><body>",
        f"<h1>{_escape(_fmt(template['header'], values))}</h1>",
        f"<p class='meta'>{_escape(_fmt(template['subheader'], values))}</p>",
    ]
    for index, question in enumerate(notebook.get("questions") or [], start=1):
        row = {**values, **_question_values(question, index)}
        parts.append(f"<h2>{_escape(_fmt(template['section_title'], row))}</h2>")
        for field in template["fields"]:
            text = _plain(question, field)
            blocks: List[str] = []
            if text:
                style = "ans mine" if field == "my_answer" else ("ans right" if field == "correct_answer" else "ans")
                blocks.append(f"<div class='{style}'>{_escape(text)}</div>")
            if field == "stem":
                for image in question.get("images") or []:
                    src = image.get("asset") or image.get("url")
                    if src:
                        blocks.append(f"<img src='{_escape(src)}' alt='题目图片'>")
            if blocks:
                parts.append(f"<div class='label'>{FIELD_LABELS[field]}</div>" + "\n".join(blocks))
    parts.append(f"<footer>{_escape(_fmt(template['footer'], values))}</footer></body></html>")
    return "\n".join(parts)


def render_docx(notebook: Dict[str, Any], template: Dict[str, Any], asset_root: str) -> bytes:
    """按模板渲染 docx。页眉页脚按要求补上（客户特意提过）。"""
    import docx
    from docx.shared import Inches

    values = dict(notebook.get("meta") or {})
    document = docx.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = _fmt(template["header"], values)
    section.footer.paragraphs[0].text = _fmt(template["footer"], values)

    document.add_heading(_fmt(template["header"], values), level=1)
    document.add_paragraph(_fmt(template["subheader"], values))
    for index, question in enumerate(notebook.get("questions") or [], start=1):
        row = {**values, **_question_values(question, index)}
        document.add_heading(_fmt(template["section_title"], row), level=2)
        for field in template["fields"]:
            text = _plain(question, field)
            if text:
                document.add_paragraph(f"{FIELD_LABELS[field]}：")
                document.add_paragraph(text)
            elif field == "stem":
                pass
            if field == "stem":
                for image in question.get("images") or []:
                    asset = image.get("asset")
                    if not asset:
                        continue
                    path = os.path.join(asset_root, asset)
                    if os.path.isfile(path):
                        document.add_picture(path, width=Inches(4))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _filename(title: str, student_name: str) -> str:
    def clean(text: str) -> str:
        return "".join(ch for ch in (text or "") if ch not in '\\/:*?"<>|').strip()

    base = "-".join(part for part in (clean(title), clean(student_name), "错题本") if part)
    return f"{base or '错题本'}.zip"


def build_notebook_zip(
    notebook: Dict[str, Any], template: Optional[Dict[str, Any]] = None
) -> Tuple[str, bytes]:
    """返回 (下载文件名, zip 内容)。包内四件套：docx / json / html / assets。"""
    merged = merge_template(template)
    values = dict(notebook.get("meta") or {})
    values["exported_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    notebook["meta"] = values

    assets = prepare_assets(notebook)
    html = render_html(notebook, merged)

    with tempfile.TemporaryDirectory() as workdir:
        for relative, payload in assets:
            target = os.path.join(workdir, relative)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with open(target, "wb") as handle:
                handle.write(payload)
        docx_bytes = render_docx(notebook, merged, workdir)

    payload = {
        "template": merged,
        "data": notebook,
        "assets": [relative for relative, _ in assets],
    }

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("notebook.json", json.dumps(payload, ensure_ascii=False, indent=2))
        archive.writestr("notebook.docx", docx_bytes)
        archive.writestr("index.html", html)
        for relative, content in assets:
            archive.writestr(relative, content)
    return _filename(values.get("title") or "", values.get("student_name") or ""), buffer.getvalue()